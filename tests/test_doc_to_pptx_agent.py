"""The doc-to-pptx agent: reading a document, placing its images, and building the deck.

`test_pptx_theme_colors.py` holds the palette to six-digit colours,
`test_doc_to_pptx_blocking.py` keeps the build off the event loop and
`test_documents.py` the paths it may use; this covers the rest. The program is
exec'd as it is at spawn and driven through a stand-in `agent` with a scripted
model. PyMuPDF and pdfplumber are stand-ins in `sys.modules`; the deck is built
by python-pptx for real and read back.
"""

import json
import sys
import tempfile
import types
from pathlib import Path
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from tests.programs import program_namespace
from wactorz.documents import DOCUMENT_DIRS_ENV
from wactorz.web.uploads import UPLOADS_DIRNAME

NS = program_namespace("doc_to_pptx_agent.py")

UPLOAD_ID = "0f3c9a5e4b2d4c1e8a7b6c5d4e3f2a1b"


def _png(path: Path, size: tuple[int, int] = (800, 600)) -> Path:
    Image.new("RGB", size, "teal").save(path, format="PNG")
    return path


def _pictures(slide: Any) -> list[Any]:
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _bullets_of(slide: Any) -> Any:
    """The text box holding a content slide's bullets, which `_outline` makes "one", "two"."""
    return next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith("one"))


@pytest.fixture(name="docs", autouse=True)
def docs_fixture(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """A document folder and a temp folder of the test's own.

    Nothing reads the real Documents, and a deck with no `output_path`, which
    goes to a fresh temp folder, is not left behind in the real one.
    """
    folder = tmp_path / "docs"
    folder.mkdir()
    monkeypatch.setenv(DOCUMENT_DIRS_ENV, str(folder))
    scratch = tmp_path / "tmp"
    scratch.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(scratch))
    return folder


def _outline(slides: int = 3, **overrides: Any) -> dict[str, Any]:
    kinds = ["title", *["content"] * (slides - 2), "closing"]
    return {
        "title": "Quarterly Review",
        "theme_colors": {"accent": "FF8800"},
        "slides": [
            {
                "index": i,
                "type": kind,
                "title": f"Slide {i}",
                "subtitle": "A subtitle" if kind != "content" else "",
                "bullets": ["one", "two"] if kind == "content" else [],
                "image_prompt": f"picture {i}",
            }
            for i, kind in enumerate(kinds)
        ],
        **overrides,
    }


class _Llm:
    def __init__(self, reply: Any) -> None:
        self.reply = reply
        self.prompts: list[str] = []

    async def chat(self, prompt: str) -> Any:
        self.prompts.append(prompt)
        return self.reply


class _Agent:
    """What the program uses of `agent`, with the image-gen agent's answers scripted."""

    def __init__(self, llm: Any = None) -> None:
        self.llm = llm
        self.logs: list[str] = []
        self.alerts: list[tuple[str, str]] = []
        self.sent: list[dict[str, Any]] = []
        self.images: dict[str, Any] = {}

    async def log(self, text: str) -> None:
        self.logs.append(text)

    async def alert(self, message: str, severity: str = "warning") -> None:
        self.alerts.append((severity, message))

    async def send_to(self, target: str, payload: dict[str, Any]) -> Any:
        assert target == "image-gen-agent"
        self.sent.append(payload)
        answer = self.images.get(payload["prompt"])
        if isinstance(answer, Exception):
            raise answer
        if answer == "made":
            _png(Path(payload["output_path"]))
            return {"image_path": payload["output_path"]}
        return answer


# ── PyMuPDF and pdfplumber ─────────────────────────────────────────────────────


class _Pixmap:
    def __init__(self, width: int, height: int, n: int = 3, alpha: int = 0) -> None:
        self.width, self.height, self.n, self.alpha = width, height, n, alpha

    def save(self, path: str) -> None:
        _png(Path(path), (self.width, self.height))


class _Page:
    def __init__(self, images: list[int]) -> None:
        self._images = images

    def get_images(self, full: bool) -> list[tuple[int]]:
        return [(xref,) for xref in self._images]


class _Doc:
    def __init__(self, pages: list[list[int]]) -> None:
        self._pages = [_Page(p) for p in pages]

    def __len__(self) -> int:
        return len(self._pages)

    def __getitem__(self, i: int) -> _Page:
        return self._pages[i]

    def close(self) -> None:
        pass


def _fitz(pages: list[list[int]], pixmaps: dict[int, Any]) -> types.ModuleType:
    """A PyMuPDF whose document has `pages` (image xrefs per page), each xref a pixmap."""
    module = types.ModuleType("fitz")
    module.csRGB = "rgb"  # pyright: ignore[reportAttributeAccessIssue]
    module.open = lambda _path: _Doc(pages)  # pyright: ignore[reportAttributeAccessIssue]

    def pixmap(source: Any, xref: Any) -> _Pixmap:
        if source == "rgb":  # the conversion of a CMYK pixmap to RGB
            return _Pixmap(xref.width, xref.height)
        found = pixmaps[xref]
        if isinstance(found, Exception):
            raise found
        return found

    module.Pixmap = pixmap  # pyright: ignore[reportAttributeAccessIssue]
    return module


class _PlumberPage:
    def __init__(self, text: str | None) -> None:
        self._text = text

    def extract_text(self) -> str | None:
        return self._text


class _Plumber:
    def __init__(self, pages: list[str | None]) -> None:
        self.pages = [_PlumberPage(p) for p in pages]

    def __enter__(self) -> "_Plumber":
        return self

    def __exit__(self, *_exc: object) -> None:
        return None


def _pdfplumber(pages: list[str | None]) -> types.ModuleType:
    module = types.ModuleType("pdfplumber")
    module.open = lambda _path: _Plumber(pages)  # pyright: ignore[reportAttributeAccessIssue]
    return module


# ── Reading ────────────────────────────────────────────────────────────────────


class TestReading:
    def test_text_in_either_encoding(self, tmp_path: Path) -> None:
        utf8 = tmp_path / "a.txt"
        utf8.write_text("café", encoding="utf-8")
        latin = tmp_path / "b.txt"
        latin.write_bytes("café".encode("latin-1"))

        assert NS["_read_document"](str(utf8)) == "café"
        assert NS["_read_document"](str(latin)) == "café"

    def test_a_pdf_is_read_page_by_page(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setitem(sys.modules, "pdfplumber", _pdfplumber([" One ", None, "Three"]))
        pdf = tmp_path / "a.PDF"
        pdf.write_bytes(b"%PDF-1.4")

        assert NS["_read_document"](str(pdf)) == "One\n\nThree"

    def test_a_pdf_is_known_by_its_bytes_not_its_name(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """An upload is stored under an id, with no extension to go by."""
        monkeypatch.setitem(sys.modules, "pdfplumber", _pdfplumber(["From the PDF"]))
        unnamed = tmp_path / UPLOAD_ID
        unnamed.write_bytes(b"%PDF-1.7")
        named = tmp_path / "notes.pdf"
        named.write_text("Plain text", encoding="utf-8")

        assert NS["_read_document"](str(unnamed)) == "From the PDF"
        assert NS["_read_document"](str(named)) == "Plain text"

    def test_text_in_no_encoding_it_knows_is_refused(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def undecodable(self: Path, encoding: str) -> str:
            raise UnicodeDecodeError(encoding, b"\xff", 0, 1, "nope")

        monkeypatch.setattr(Path, "read_text", undecodable)

        with pytest.raises(ValueError, match="Cannot decode file"):
            NS["_read_txt"](str(tmp_path / "a.txt"))


class TestImages:
    def test_large_images_are_saved_once_and_small_or_broken_ones_skipped(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        fitz = _fitz(
            pages=[[1, 2], [1, 3, 4]],
            pixmaps={
                1: _Pixmap(800, 600),
                2: _Pixmap(50, 50),  # a logo
                3: _Pixmap(900, 700, n=5, alpha=1),  # CMYK with alpha
                4: RuntimeError("unsupported"),
            },
        )
        monkeypatch.setitem(sys.modules, "fitz", fitz)

        images = NS["_extract_pdf_images"]("doc.pdf", str(tmp_path))

        assert [(i["page"], i["width"]) for i in images] == [(0, 800), (1, 900)]
        assert all(Path(i["path"]).exists() for i in images)

    def test_title_and_closing_take_the_largest_and_content_its_own_pages(self) -> None:
        images = [
            {"path": "big", "page": 9, "width": 900, "height": 900},
            {"path": "mid", "page": 3, "width": 600, "height": 400},
            {"path": "early", "page": 0, "width": 300, "height": 200},
        ]
        slides = [
            {"index": 0, "type": "title"},
            {"index": 1},
            {"index": 2},
            {"index": 3, "type": "closing"},
        ]

        assignment = NS["_assign_images_to_slides"](images, slides, 10)

        # Four slides over ten pages: the first content slide covers pages 2-5, the
        # second 5-7; the closing slide takes the largest image left.
        assert assignment == {0: "big", 1: "mid", 2: None, 3: "early"}

    def test_nothing_to_place_places_nothing(self) -> None:
        assert NS["_assign_images_to_slides"]([], [{"index": 0}], 3) == {}

    async def test_slides_without_a_picture_ask_for_one(self, tmp_path: Path) -> None:
        agent = _Agent()
        agent.images = {
            "made one": "made",
            "refused": {"error": "quota"},
            "broke": RuntimeError("down"),
        }
        slides = [
            {"index": 0, "image_prompt": "made one"},
            {"index": 1, "image_prompt": "refused"},
            {"index": 2, "image_prompt": "broke"},
            {"index": 3},
            {"index": 4, "image_prompt": "unused"},
        ]
        assignment: dict[int, Any] = {0: None, 1: None, 2: None, 3: None, 4: "pdf.png"}

        generated = await NS["_nim_generate_missing"](agent, slides, assignment, str(tmp_path))

        assert generated == 1
        assert assignment[0] == str(tmp_path / "nim_img_0.png")
        assert assignment[1] is assignment[2] is assignment[3] is None
        assert assignment[4] == "pdf.png"
        assert {p["prompt"] for p in agent.sent} == {"made one", "refused", "broke"}
        assert any("slide 2 failed: down" in line for line in agent.logs)

    async def test_with_every_slide_pictured_nothing_is_asked(self, tmp_path: Path) -> None:
        agent = _Agent()

        generated = await NS["_nim_generate_missing"](
            agent, [{"index": 0}], {0: "a.png"}, str(tmp_path)
        )

        assert generated == 0
        assert agent.sent == []


class TestTheOutline:
    @pytest.mark.parametrize(
        "reply",
        [
            json.dumps(_outline()),
            "```json\n" + json.dumps(_outline()) + "\n```",
            "```\n" + json.dumps(_outline()) + "\n```",
        ],
    )
    async def test_it_is_read_with_or_without_fences(self, reply: str) -> None:
        llm = _Llm(reply)

        outline = await NS["_extract_outline"](_Agent(llm), "x" * 5000, 3)

        assert outline["title"] == "Quarterly Review"
        assert "Produce exactly 3 slides" in llm.prompts[0]
        assert "x" * 4001 not in llm.prompts[0]


class TestTheDeck:
    @staticmethod
    def _build(tmp_path: Path, outline: dict[str, Any], pictured: dict[int, Any]) -> Any:
        out = tmp_path / "deck.pptx"
        NS["_build_deck"](outline, pictured, str(out))
        return Presentation(str(out))

    def test_each_kind_of_slide_with_and_without_its_picture(self, tmp_path: Path) -> None:
        picture = str(_png(tmp_path / "p.png"))
        pictured = {0: picture, 1: picture, 2: None, 3: picture}

        deck = self._build(tmp_path, _outline(4), pictured)

        slides = list(deck.slides)
        assert [len(_pictures(s)) for s in slides] == [1, 1, 0, 1]
        assert (deck.slide_width, deck.slide_height) == (9144000, 5143500)
        assert deck.core_properties.title == "Quarterly Review"

    def test_text_sits_beside_a_picture_and_across_a_slide_without_one(
        self, tmp_path: Path
    ) -> None:
        picture = str(_png(tmp_path / "p.png"))

        deck = self._build(tmp_path, _outline(4), {1: picture, 2: None})

        widths = [_bullets_of(slide).width for slide in list(deck.slides)[1:3]]
        assert widths == [NS["_emu"](4.7), NS["_emu"](9.3)]

    def test_bullets_are_bullets(self, tmp_path: Path) -> None:
        deck = self._build(tmp_path, _outline(3), {})

        assert _bullets_of(list(deck.slides)[1]).element.xml.count("<a:buChar") == 2

    def test_a_picture_is_cropped_to_its_box_not_stretched(self, tmp_path: Path) -> None:
        tall = str(_png(tmp_path / "tall.png", (400, 900)))

        deck = self._build(tmp_path, _outline(3), {1: tall})

        [picture] = _pictures(list(deck.slides)[1])
        assert picture.crop_left == picture.crop_right == 0
        assert picture.crop_top == picture.crop_bottom > 0

    def test_the_closing_picture_is_faded(self, tmp_path: Path) -> None:
        picture = str(_png(tmp_path / "p.png"))

        deck = self._build(tmp_path, _outline(3), {2: picture})

        [faded] = _pictures(list(deck.slides)[2])
        assert '<a:alphaModFix amt="45000"/>' in faded.element.xml

    def test_a_picture_that_cannot_be_read_leaves_a_text_slide(self, tmp_path: Path) -> None:
        broken = tmp_path / "broken.png"
        broken.write_bytes(b"not a picture")

        deck = self._build(tmp_path, _outline(3), {1: str(broken)})

        assert _pictures(list(deck.slides)[1]) == []


# ── The whole conversion ───────────────────────────────────────────────────────


def _txt(tmp_path: Path) -> Path:
    path = tmp_path / "docs" / "report.txt"
    path.write_text("Revenue grew. Costs fell. " * 50, encoding="utf-8")
    return path


class TestConverting:
    async def test_a_text_document_becomes_a_deck(self, tmp_path: Path) -> None:
        agent = _Agent(_Llm(json.dumps(_outline(3))))
        agent.images = {"picture 0": "made", "picture 1": "made", "picture 2": "made"}
        out = tmp_path / "docs" / "out" / "deck.pptx"
        out.parent.mkdir()

        result = await NS["handle_task"](
            agent, {"file_path": str(_txt(tmp_path)), "output_path": str(out), "slide_count": 3}
        )

        assert result["error"] is None
        assert result["pptx_path"] == str(out)
        assert (result["slide_count"], result["images_generated"]) == (3, 3)
        deck = Presentation(str(out))
        assert [len(_pictures(s)) for s in deck.slides] == [1, 1, 1]

    async def test_a_pdf_brings_its_own_pictures(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        pdf = tmp_path / "docs" / "report.pdf"
        pdf.write_bytes(b"%PDF-1.4")
        monkeypatch.setitem(sys.modules, "pdfplumber", _pdfplumber(["Revenue grew."]))
        monkeypatch.setitem(
            sys.modules,
            "fitz",
            _fitz(
                [[1], [2], [3]], {1: _Pixmap(800, 600), 2: _Pixmap(700, 500), 3: _Pixmap(600, 400)}
            ),
        )
        agent = _Agent(_Llm(json.dumps(_outline(3))))

        result = await NS["handle_task"](agent, {"file_path": str(pdf), "nim_fallback": False})

        assert result["images_extracted"] == 3
        assert result["images_generated"] == 0
        assert agent.sent == []
        assert Path(result["pptx_path"]).exists()

    async def test_without_pymupdf_a_pdf_is_converted_without_its_pictures(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        pdf = tmp_path / "docs" / "report.pdf"
        pdf.write_bytes(b"%PDF-1.4")
        monkeypatch.setitem(sys.modules, "pdfplumber", _pdfplumber(["Revenue grew."]))
        monkeypatch.setitem(sys.modules, "fitz", None)
        agent = _Agent(_Llm(json.dumps(_outline(3))))

        result = await NS["handle_task"](agent, {"file_path": str(pdf), "nim_fallback": False})

        assert result["error"] is None
        assert result["images_extracted"] == 0
        assert any("PyMuPDF not installed" in line for line in agent.logs)
        assert any("text-only" in line for line in agent.logs)

    async def test_the_request_may_arrive_as_text(self, tmp_path: Path) -> None:
        agent = _Agent(_Llm(json.dumps(_outline(3))))
        request = json.dumps({"file_path": str(_txt(tmp_path)), "nim_fallback": False})

        as_string = await NS["handle_task"](agent, request)
        in_text = await NS["handle_task"](agent, {"message": "{oops", "text": request})

        assert as_string["error"] is None
        assert in_text["error"] is None

    @pytest.mark.parametrize("request_", [{}, "convert something"])
    async def test_a_request_naming_no_file_is_said(self, request_: Any) -> None:
        result = await NS["handle_task"](_Agent(), request_)

        assert result["error"].startswith("Name the file")

    async def test_a_missing_file_is_said(self, docs: Path) -> None:
        result = await NS["handle_task"](_Agent(), {"file_path": str(docs / "gone.txt")})

        assert result["error"].startswith("File not found")

    async def test_a_file_outside_the_document_folders_is_not_read(self, tmp_path: Path) -> None:
        secret = tmp_path / "secret.txt"
        secret.write_text("the deploy key", encoding="utf-8")
        llm = _Llm(json.dumps(_outline(3)))

        result = await NS["handle_task"](_Agent(llm), {"file_path": str(secret)})

        assert "outside the document folders" in result["error"]
        assert llm.prompts == []

    @pytest.mark.parametrize(
        ("where", "said"),
        [
            ("elsewhere.pptx", "outside the document folders"),
            ("docs/.bashrc", "must end in .pptx"),
            ("state/agent.pptx", "outside the document folders"),
        ],
    )
    async def test_a_deck_is_not_written_where_it_may_not_go(
        self, where: str, said: str, tmp_path: Path
    ) -> None:
        llm = _Llm(json.dumps(_outline(3)))

        result = await NS["handle_task"](
            _Agent(llm), {"file_path": str(_txt(tmp_path)), "output_path": str(tmp_path / where)}
        )

        assert said in result["error"]
        assert llm.prompts == []
        assert not (tmp_path / where).exists()

    async def test_an_attached_file_is_converted_by_its_id(self, tmp_path: Path) -> None:
        stored = tmp_path / "state" / UPLOADS_DIRNAME / UPLOAD_ID
        stored.parent.mkdir(parents=True)
        stored.write_text("Revenue grew. " * 20, encoding="utf-8")
        agent = _Agent(_Llm(json.dumps(_outline(3))))

        result = await NS["handle_task"](agent, {"file_id": UPLOAD_ID, "nim_fallback": False})

        assert result["error"] is None
        assert len(Presentation(result["pptx_path"]).slides) == 3

    async def test_an_empty_document_is_said(self, tmp_path: Path) -> None:
        empty = tmp_path / "docs" / "empty.txt"
        empty.write_text("   ", encoding="utf-8")

        result = await NS["handle_task"](_Agent(), {"file_path": str(empty)})

        assert result["error"] == "Document appears empty or unreadable"

    @pytest.mark.parametrize(
        ("reply", "said"),
        [
            ("Here is your outline!", "LLM outline JSON parse failed"),
            ("[1, 2]", "doc-to-pptx failed"),
        ],
    )
    async def test_an_outline_that_cannot_be_used_is_said(
        self, reply: str, said: str, tmp_path: Path
    ) -> None:
        agent = _Agent(_Llm(reply))

        result = await NS["handle_task"](agent, {"file_path": str(_txt(tmp_path))})

        assert result["error"].startswith(said)
        assert agent.alerts[0][0] == "error"

    @pytest.mark.parametrize(
        ("failure", "said"),
        [
            (RuntimeError("disk full"), "Building the deck failed: disk full"),
            (ImportError("pptx"), "python-pptx is not installed: pip install python-pptx"),
        ],
    )
    async def test_a_failed_build_is_said(
        self, failure: Exception, said: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def fails(*_args: Any) -> None:
            raise failure

        monkeypatch.setitem(NS, "_build_deck", fails)
        agent = _Agent(_Llm(json.dumps(_outline(3))))

        result = await NS["handle_task"](
            agent, {"file_path": str(_txt(tmp_path)), "nim_fallback": False}
        )

        assert result["error"] == said
        assert result["pptx_path"] is None
        assert result["slide_count"] == 3

    @pytest.mark.parametrize("field", ["slide_count", "min_img_width", "min_img_height"])
    async def test_a_number_that_is_not_one_is_said(self, field: str, tmp_path: Path) -> None:
        llm = _Llm(json.dumps(_outline(3)))

        result = await NS["handle_task"](
            _Agent(llm), {"file_path": str(_txt(tmp_path)), field: "many"}
        )

        assert result["error"] == f"{field} must be a whole number"
        assert llm.prompts == []

    async def test_the_working_files_are_removed_and_the_deck_kept(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        scratch = tmp_path / "tmp"
        agent = _Agent(_Llm(json.dumps(_outline(3))))
        agent.images = {"picture 0": "made"}

        done = await NS["handle_task"](agent, {"file_path": str(_txt(tmp_path))})

        def fails(*_args: Any) -> None:
            raise RuntimeError("disk full")

        monkeypatch.setitem(NS, "_build_deck", fails)
        failed = await NS["handle_task"](agent, {"file_path": str(_txt(tmp_path))})

        assert Path(done["pptx_path"]).exists()
        assert failed["pptx_path"] is None
        # Only the decks' own folders are left: the generated pictures are gone.
        assert all(p.name.startswith("wactorz-pptx-") for p in scratch.iterdir())


class TestStarting:
    _MODULES = ("pptx", "fitz", "pdfplumber", "PIL")

    async def test_a_ready_machine_is_said_once(self, monkeypatch: pytest.MonkeyPatch) -> None:
        for name in self._MODULES:
            monkeypatch.setitem(sys.modules, name, types.ModuleType(name))
        agent = _Agent()

        await NS["setup"](agent)

        assert agent.alerts == []
        assert "ready" in agent.logs[0]

    async def test_what_is_missing_is_alerted(self, monkeypatch: pytest.MonkeyPatch) -> None:
        for name in self._MODULES:
            monkeypatch.setitem(sys.modules, name, None)
        agent = _Agent()

        await NS["setup"](agent)

        [(severity, message)] = agent.alerts
        assert severity == "warning"
        assert "pip install python-pptx" in message
        assert "pip install pymupdf" in message

    async def test_there_is_nothing_to_poll(self, monkeypatch: pytest.MonkeyPatch) -> None:
        slept: list[float] = []

        async def sleep(seconds: float) -> None:
            slept.append(seconds)

        monkeypatch.setitem(NS, "asyncio", types.SimpleNamespace(sleep=sleep))

        await NS["process"](_Agent())

        assert slept == [3600]
